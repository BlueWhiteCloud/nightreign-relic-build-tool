"""
从「局外遗物词条详细数据.xlsx」抽出增伤类词条的数值，生成程序可直接读的 CSV。

为什么要抽成 CSV：
    程序运行时不该依赖 openpyxl（打包 exe 会多一个依赖），而且这张表是人工维护的，
    抽一次、核对一次即可；以后表更新了，重跑这个脚本就行。

用法：
    python tools/build_attack_values.py

产物：
    Resources/Param/AttackValues.csv
    列：effect_id, category_id, kind, value_percent, name, note
        value_percent  = 解释里第一个「数字%」的数值（如 12 表示 +12%）
        note           = 原始解释文本，方便人工核对
"""
import csv
import re
import sys
from pathlib import Path

WORKING_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(WORKING_DIR))

import openpyxl  # noqa: E402

from recommender.effect_table import EffectTable  # noqa: E402

XLSX = WORKING_DIR.parent / "局外遗物词条详细数据.xlsx"
OUT = WORKING_DIR / "Resources" / "Param" / "AttackValues.csv"

# 收录哪几类（用户口径）：
#     攻击力      —— 泛用增伤（属性/武器类别/战技强化…）
#     魔法/祷告   —— 流派增伤（「强化卡利亚剑的魔法」「强化野兽的祷告」…统一 +12%）
WANTED_KINDS = ("攻击力", "魔法/祷告")

# 下面这些整类收会带进一堆无关词条，所以按「类型 + 名字关键词」单独收：
#     「装备三把以上类别为短剑/直剑/…/弓的武器，能提升攻击力」（+20%，弓 +10%）
#     「对陷入中毒/猩红腐败/冻伤的敌人，能强化攻击（＋１／＋２）」（+10% / +16% / +20%）
#     「周围陷入催眠/发狂时，提升攻击力（＋１）」「周围人物陷入中毒、腐败时，能提升攻击力」
#     「提升近战攻击力」（+6%）「提升战技攻击力」（+15%）—— 安定的遗志那两件专属遗物
EXTRA_RULES = (
    ("仅限特定武器", "能提升攻击力"),
    ("行动", "能强化攻击"),
    ("行动", "提升攻击力"),
    ("专属遗物", "提升近战攻击力"),
    ("专属遗物", "提升战技攻击力"),
)

PERCENT_RE = re.compile(r"(\d+(?:\.\d+)?)\s*%")


def _cell_text(value) -> str:
    return "" if value is None else str(value).replace("\n", " ").strip()


def main() -> int:
    if not XLSX.exists():
        print(f"找不到表格：{XLSX}")
        return 1

    effect_table = EffectTable()
    workbook = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)

    rows = []
    skipped = []
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            effect_id, name, note, kind, _stack = row
            if effect_id is None:
                continue
            try:
                effect_id = int(effect_id)
            except (TypeError, ValueError):
                continue
            kind_text = _cell_text(kind)
            name_text = _cell_text(name)
            wanted = kind_text in WANTED_KINDS or any(
                kind_text == kind and keyword in name_text
                for kind, keyword in EXTRA_RULES)
            if not wanted:
                continue

            meta = effect_table.effect(effect_id)
            if meta is None:
                skipped.append((effect_id, name_text, "程序词条表里没有这个 effect_id"))
                continue

            note_text = _cell_text(note)
            match = PERCENT_RE.search(note_text)
            if match is None:
                skipped.append((effect_id, name_text, f"解释里没有百分比：{note_text}"))
                continue

            rows.append({
                "effect_id": effect_id,
                "category_id": meta.category_id,
                "kind": kind_text,
                "value_percent": match.group(1),
                "name": name_text,
                "note": note_text,
            })

    rows.sort(key=lambda r: r["effect_id"])
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["effect_id", "category_id", "kind",
                                               "value_percent", "name", "note"])
        writer.writeheader()
        writer.writerows(rows)

    print(f"已写出 {len(rows)} 条增伤词条（{'、'.join(WANTED_KINDS)}"
          f" ＋ 按名字补收的 {'、'.join(f'{k}·{kw}' for k, kw in EXTRA_RULES)}）→ {OUT}")
    print(f"词条编号去重后 {len({r['category_id'] for r in rows})} 个")
    if skipped:
        print(f"\n跳过 {len(skipped)} 条（需人工确认）：")
        for effect_id, name, why in skipped:
            print(f"  {effect_id} {name} —— {why}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
